#!/usr/bin/env python3
"""Converte l'export Excel del calendario FIPAV in data/partite.json.

Uso:
    python scripts/converti.py [file.xlsx] [squadra] [nome mostrato]

Default: data/calendario_gare.xlsx, squadra "MANDELLO".
Tiene solo le partite in cui la squadra gioca (in casa o in trasferta).
"""
import json
import sys
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "calendario_gare.xlsx"
SQUADRA = (sys.argv[2] if len(sys.argv) > 2 else "MANDELLO").upper()
NOME = sys.argv[3] if len(sys.argv) > 3 else "Mandello"   # nome mostrato nella pagina
OUT = ROOT / "data" / "partite.json"
IMPIANTI = ROOT / "data" / "impianti.json"   # correzioni manuali degli indirizzi


def testo(v):
    return str(v).strip() if v is not None else ""


def leggi_data(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = testo(v)
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Data non riconosciuta: {v!r}")


def leggi_ora(v):
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, time):
        return v
    s = testo(v).replace(".", ":")
    return datetime.strptime(s, "%H:%M").time()


def main():
    ws = load_workbook(XLSX, data_only=True).active
    righe = list(ws.iter_rows(values_only=True))
    intest = [testo(c).lower() for c in righe[0]]

    def col(nome):
        try:
            return intest.index(nome)
        except ValueError:
            raise SystemExit(f"Colonna '{nome}' non trovata. Colonne: {intest}")

    c = {k: col(k) for k in [
        "giornata", "a/r", "data", "ora", "casa", "trasferta",
        "nome impianto", "indirizzo", "citta", "frazione", "provincia",
    ]}
    c_serie = intest.index("serie") if "serie" in intest else None
    c_girone = intest.index("girone") if "girone" in intest else None

    correzioni = {}
    if IMPIANTI.exists():
        for k, v in json.loads(IMPIANTI.read_text(encoding="utf-8")).items():
            if not k.startswith("_"):
                correzioni[k.upper()] = v

    partite = []
    for r in righe[1:]:
        if not any(x is not None for x in r):
            continue
        casa, trasf = testo(r[c["casa"]]), testo(r[c["trasferta"]])
        in_casa = SQUADRA in casa.upper()
        if not in_casa and SQUADRA not in trasf.upper():
            continue
        g = leggi_data(r[c["data"]])
        o = leggi_ora(r[c["ora"]])
        inizio = datetime.combine(g, o)
        citta = testo(r[c["citta"]]).title()
        frazione = testo(r[c["frazione"]]).title()
        prov = testo(r[c["provincia"]]).upper()
        via = testo(r[c["indirizzo"]]).title()
        luogo = ", ".join(x for x in [via, frazione, f"{citta} ({prov})" if prov else citta] if x)
        impianto = testo(r[c["nome impianto"]]).title()
        q_maps = f"{impianto}, {via}, {citta} {prov}"
        q_nav = f"{via}, {citta}, {prov}, Italia" if via else f"{citta}, Italia"
        fix = correzioni.get(f"{impianto} - {citta}".upper())
        if fix:
            luogo = fix.get("indirizzo", luogo)
            q_nav = fix.get("indirizzo", q_nav)
            q_maps = f"{impianto}, {fix['indirizzo']}" if "indirizzo" in fix else q_maps
            if "lat" in fix and "lon" in fix:
                q_nav = q_maps = f"{fix['lat']},{fix['lon']}"
        partite.append({
            "giornata": int(r[c["giornata"]]) if r[c["giornata"]] is not None else None,
            "andata_ritorno": testo(r[c["a/r"]]),
            "inizio": inizio.strftime("%Y-%m-%dT%H:%M"),
            "in_casa": in_casa,
            "casa": casa,
            "trasferta": trasf,
            "avversario": trasf if in_casa else casa,
            "impianto": impianto,
            "indirizzo": luogo,
            "cerca_maps": q_maps,
            "cerca_waze": q_nav,
        })

    if not partite:
        raise SystemExit(f"Nessuna partita trovata per '{SQUADRA}'")

    partite.sort(key=lambda p: p["inizio"])
    primo = righe[1]
    out = {
        "squadra": SQUADRA,
        "nome": NOME,
        "campionato": {"DM": "Serie D Maschile", "DF": "Serie D Femminile", "CM": "Serie C Maschile", "CF": "Serie C Femminile"}.get(testo(primo[c_serie]) if c_serie is not None else "", testo(primo[c_serie]) if c_serie is not None else "Serie D"),
        "girone": testo(primo[c_girone]) if c_girone is not None else "",
        "aggiornato": datetime.now().strftime("%Y-%m-%dT%H:%M"),
        "ritrovo_minuti_prima": 75,
        "partite": partite,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(partite)} partite di {SQUADRA} scritte in {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
